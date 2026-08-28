#!/usr/bin/env python3
"""
Preenche o template de Caracteristicas de Risco (Empresarial), N massas = N linhas.

Uso:
    python3 preencher_risco.py entrada_risco.json saida.xlsx template.xlsx atividades.txt

Formato do JSON:
    {"massas": [
        {"tipo_construcao": "Superior",
         "objeto_segurado": "Predio e Conteudo",
         "assistencia": "Assistencia Empresarial Plano Basico",
         "atividade": "Academias",
         "valor_em_risco": 1000000}
    ]}

Regras:
  - Colunas lidas pelo NOME do cabecalho, nunca por posicao.
  - Campos residenciais (CHK Residencial Beneficios Essenciais, CBO Tipo de
    Residencia) recebem sempre <IGNORE> no empresarial.
  - Valor em Risco sai no formato BR do template: 1.000.000,00
  - Valor fora da lista permitida ou atividade fora do catalogo = erro.
"""

import json
import sys

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from preencher_cobertura import ABA, IGNORAR, LINHA_CABECALHO, LINHA_DADOS, normalizar

OBJETO_SEGURADO = ["Prédio e Conteúdo", "Prédio", "Conteúdo"]
TIPO_RESIDENCIA = [
    "Apartamento Desocupado",
    "Apartamento Habitual",
    "Apartamento Veraneio",
    "Casa Desocupada",
    "Casa em Condomínio Fechado Desocupada",
    "Casa em Condomínio Fechado Habitual",
    "Casa em Condomínio Fechado Veraneio",
    "Casa Habitual",
    "Casa Veraneio",
]

# Mesmo template, dois ramos. "Selecione um item" e placeholder de combo,
# nunca sai nas massas. Colunas fora do ramo recebem <IGNORE>.
RAMOS = {
    "empresarial": {
        "tipo_construcao": ["Superior", "Sólida", "Mista", "Inferior"],
        "objeto_segurado": OBJETO_SEGURADO,
        "assistencia": [
            "Assistência Empresarial Plano Superior",
            "Assistência Empresarial Plano Intermediário",
            "Assistência Empresarial Plano Básico",
            "Não Contratado",
        ],
        "tipo_residencia": None,   # coluna F fica <IGNORE>
        "atividade": True,         # coluna E vem do atividades.txt
    },
    "residencial": {
        "tipo_construcao": ["Superior / Sólida", "Mista", "Inferior"],
        "objeto_segurado": OBJETO_SEGURADO,
        "assistencia": [
            "Assistência Residencial Plano Superior",
            "Assistência Residencial Plano Intermediário",
            "Assistência Residencial Plano Básico",
            "Não Contratado",
        ],
        "tipo_residencia": TIPO_RESIDENCIA,
        "atividade": False,        # coluna E fica <IGNORE>
    },
}

COL_CONSTRUCAO = "CBO Tipo de Construção"
COL_OBJETO = "CBO Objeto Segurado"
COL_ASSIST = "CBO Assistência 24h"
COL_BENEF = "CHK Residencial Benefícios Essenciais"
COL_ATIV = "TXT Atividade"
COL_TIPO_RES = "CBO Tipo de Residência"
COL_VALOR = "TXT Valor em Risco - Danos Materiais"
COLUNAS_ESPERADAS = [
    COL_CONSTRUCAO, COL_OBJETO, COL_ASSIST, COL_BENEF,
    COL_ATIV, COL_TIPO_RES, COL_VALOR,
]


def mapear(ws):
    colunas = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=LINHA_CABECALHO, column=c).value
        if h:
            colunas[str(h).strip()] = get_column_letter(c)
    faltando = [h for h in COLUNAS_ESPERADAS if h not in colunas]
    if faltando:
        raise ValueError(f"Cabecalho sem as colunas esperadas: {faltando}")
    return colunas


def escolher(valor, permitidos, campo, i):
    alvo = normalizar(valor)
    for p in permitidos:
        if normalizar(p) == alvo:
            return p
    raise ValueError(
        f"massa {i}: {campo} = {valor!r} nao esta na lista. Validos: {permitidos}"
    )


def formatar_reais(v):
    """1000000 -> '1.000.000,00'."""
    if isinstance(v, str):
        v = v.replace(".", "").replace(",", ".")
    inteiro = f"{int(round(float(v))):,}".replace(",", ".")
    return f"{inteiro},00"


def preencher(dados, caminho_template, caminho_saida, atividades=None):
    ramo = dados.get("ramo", "empresarial").strip().lower()
    if ramo not in RAMOS:
        raise ValueError(f"ramo invalido: {ramo!r}. Use 'empresarial' ou 'residencial'.")
    regras = RAMOS[ramo]

    wb = openpyxl.load_workbook(caminho_template)
    ws = wb[ABA]
    colunas = mapear(ws)
    idx_ativ = {normalizar(a): a for a in (atividades or [])}

    resumo = []
    for i, massa in enumerate(dados.get("massas", []), start=1):
        linha = LINHA_DADOS + i - 1
        valores = {
            COL_CONSTRUCAO: escolher(massa["tipo_construcao"], regras["tipo_construcao"], "tipo_construcao", i),
            COL_OBJETO: escolher(massa["objeto_segurado"], regras["objeto_segurado"], "objeto_segurado", i),
            COL_ASSIST: escolher(massa["assistencia"], regras["assistencia"], "assistencia", i),
            COL_VALOR: formatar_reais(massa["valor_em_risco"]),
            COL_BENEF: IGNORAR,
        }

        if regras["atividade"]:
            ativ = idx_ativ.get(normalizar(massa["atividade"]))
            if ativ is None:
                raise ValueError(f"massa {i}: atividade fora do catalogo: {massa['atividade']!r}")
            valores[COL_ATIV] = ativ
        else:
            valores[COL_ATIV] = IGNORAR

        if regras["tipo_residencia"]:
            valores[COL_TIPO_RES] = escolher(
                massa["tipo_residencia"], regras["tipo_residencia"], "tipo_residencia", i
            )
        else:
            valores[COL_TIPO_RES] = IGNORAR

        for h, v in valores.items():
            ws[f"{colunas[h]}{linha}"] = v
        resumo.append((linha, valores))

    for linha_ws in ws.iter_rows(min_row=LINHA_DADOS, max_row=ws.max_row):
        for celula in linha_ws:
            celula.alignment = Alignment()

    wb.save(caminho_saida)
    return ramo, resumo


def main():
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(1)
    entrada, saida, template = sys.argv[1:4]
    cat = sys.argv[4] if len(sys.argv) > 4 else None

    with open(entrada, encoding="utf-8") as f:
        dados = json.load(f)
    atividades = None
    if cat:
        with open(cat, encoding="utf-8") as f:
            atividades = [l.strip() for l in f if l.strip()]

    ramo, resumo = preencher(dados, template, saida, atividades)
    print(f"{saida}: {len(resumo)} massas ({ramo}).")
    for linha, v in resumo:
        print(
            f"  linha {linha}: {v[COL_CONSTRUCAO]} | {v[COL_OBJETO]} | {v[COL_ASSIST]} | "
            f"ativ={v[COL_ATIV]} | res={v[COL_TIPO_RES]} | {v[COL_VALOR]}"
        )


if __name__ == "__main__":
    main()
